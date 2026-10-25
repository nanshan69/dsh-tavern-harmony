#!/usr/bin/env python3
"""CSV ⇄ XLSX 互转与 Markdown 表格导出（依赖 openpyxl，纯 Python 包）。

用法：
  python scripts/xlsx_tools.py csv2xlsx evidence/evidence.csv --out evidence/evidence.xlsx
  python scripts/xlsx_tools.py csv2xlsx a.csv b.csv --out 合集.xlsx        # 每个 CSV 一个 sheet
  python scripts/xlsx_tools.py xlsx2csv 输入.xlsx --out 导出.csv [--sheet Sheet1]
  python scripts/xlsx_tools.py xlsx2md  输入.xlsx [--sheet Sheet1]         # 打印 Markdown 表格
  python scripts/xlsx_tools.py sheets   输入.xlsx
"""
import argparse
import csv
import os
import sys

sys.path.insert(0, os.environ.get("DSH_TOOL_DIR") or os.path.dirname(os.path.abspath(globals().get("__file__", "run_tool.py"))))
from _util import ensure_parent, md_table, resolve_path  # noqa: E402


def _openpyxl():
    try:
        import openpyxl  # noqa: WPS433
        return openpyxl
    except ImportError:
        print("[缺少依赖] openpyxl 未安装。先运行：\n  python scripts/deps.py install\n"
              "（需要联网；装到内嵌 CPython 的 site-packages）", file=sys.stderr)
        raise SystemExit(3)


def _disp_width(value) -> int:
    """单元格显示宽度：中文/全角按 2，取多行中最长一行。"""
    import unicodedata
    lines = str(value).split("\n")
    return max(sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in ln)
               for ln in lines)


def read_csv_rows(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def csv2xlsx(csv_paths, out, sheet_names=None) -> None:
    openpyxl = _openpyxl()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for idx, cp in enumerate(csv_paths):
        rows = read_csv_rows(cp)
        name = (sheet_names[idx] if sheet_names and idx < len(sheet_names)
                else cp.stem)[:31] or f"Sheet{idx + 1}"
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        if rows:
            from openpyxl.styles import Alignment, Font, PatternFill
            head_fill = PatternFill("solid", fgColor="D9E2F3")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            # 冻结表头 + 首列（首列通常是 PMID/编号）；只有一列时只冻结表头
            ws.freeze_panes = "B2" if ws.max_column > 1 else "A2"
            ws.auto_filter.ref = ws.dimensions
            wrap = Alignment(wrap_text=True, vertical="top")
            for col in ws.columns:
                # 中文按 2 个字符宽；取每格最长一行；上限 50，超出部分自动换行
                width = max((_disp_width(c.value) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(8, width * 1.1 + 2), 50)
                for c in col[1:]:
                    if c.value is not None and (_disp_width(c.value) > 45 or "\n" in str(c.value)):
                        c.alignment = wrap
    ensure_parent(out)
    wb.save(out)


def xlsx2csv(xlsx_path, out, sheet=None) -> None:
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    ensure_parent(out)
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])


def xlsx2md(xlsx_path, sheet=None, limit=200) -> None:
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    for name in ([sheet] if sheet else wb.sheetnames):
        ws = wb[name]
        rows = [[("" if v is None else str(v)) for v in row]
                for row in ws.iter_rows(values_only=True)][:limit]
        if not rows:
            continue
        print(f"\n## {name}\n")
        print(md_table(rows[0], rows[1:]))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="CSV/XLSX 互转（openpyxl）")
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("csv2xlsx")
    c.add_argument("csv", nargs="+")
    c.add_argument("--out", required=True)
    c.add_argument("--sheet", nargs="*")
    x = sub.add_parser("xlsx2csv")
    x.add_argument("xlsx")
    x.add_argument("--out", required=True)
    x.add_argument("--sheet")
    m = sub.add_parser("xlsx2md")
    m.add_argument("xlsx")
    m.add_argument("--sheet")
    m.add_argument("--limit", type=int, default=200)
    s = sub.add_parser("sheets")
    s.add_argument("xlsx")
    args = ap.parse_args(argv)
    if args.cmd == "csv2xlsx":
        csvs = [resolve_path(p) for p in args.csv]
        out = resolve_path(args.out)
        csv2xlsx(csvs, out, args.sheet)
        print(f"[完成] {len(csvs)} 个 CSV → {out}")
    elif args.cmd == "xlsx2csv":
        xlsx2csv(resolve_path(args.xlsx), resolve_path(args.out), args.sheet)
        print(f"[完成] {args.xlsx} → {args.out}")
    elif args.cmd == "xlsx2md":
        xlsx2md(resolve_path(args.xlsx), args.sheet, args.limit)
    else:
        openpyxl = _openpyxl()
        wb = openpyxl.load_workbook(resolve_path(args.xlsx), read_only=True)
        print(f"{args.xlsx} 的 sheet：")
        for n in wb.sheetnames:
            print(f"  - {n}（{wb[n].max_row} 行 × {wb[n].max_column} 列）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
